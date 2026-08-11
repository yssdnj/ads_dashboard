from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .constants import EXPECTED_KEYWORD_COUNTS, RELEVANCE_LEVELS


SOURCE_TO_LEVEL = {"高相关": "high", "中相关": "mid", "低相关": "low"}

SHEET_NAME = "关键词列表"
KEYWORD_COLUMN = "关键词 (数据来源于西柚洞察)"
TRANSLATION_COLUMN = "翻译"
RELEVANCE_COLUMN = "相关性"
RELEVANCE_WEIGHT_COLUMN = "类目相关度"

REQUIRED_COLUMNS = (
    KEYWORD_COLUMN,
    TRANSLATION_COLUMN,
    "关键词排名",
    "月搜索量",
    RELEVANCE_COLUMN,
    RELEVANCE_WEIGHT_COLUMN,
    "类目相关搜索量",
    "CPC建议竞价($)",
    "建议竞价范围($)",
    "点击转化率(均值)",
    "竞争难度",
    "自然位滚动率",
)


class CatalogValidationError(ValueError):
    """Raised when the Xydc keyword catalog does not meet its contract."""


@dataclass(frozen=True)
class KeywordCatalogRow:
    keyword: str
    translation: str | None
    relevance_level: str
    relevance_weight: float
    keyword_rank: int | None
    search_volume: float | None
    category_search_volume: float | None
    cpc: float | None
    cpc_range: str | None
    click_conversion_rate: float | None
    competitive_difficulty: float | None
    organic_scroll_rate: float | None


def normalize_keyword(value: object) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _optional_text(value: object) -> str | None:
    text = str(value or "").strip()
    return text or None


def _optional_float(value: object, *, column: str, keyword: str) -> float | None:
    if isinstance(value, bool):
        raise CatalogValidationError(f"invalid {column} for {keyword}: {value}")
    if value is None or value == "":
        return None
    try:
        numeric_value = float(value)
    except (TypeError, ValueError) as exc:
        raise CatalogValidationError(f"invalid {column} for {keyword}: {value}") from exc
    if not math.isfinite(numeric_value):
        raise CatalogValidationError(f"invalid {column} for {keyword}: {value}")
    return numeric_value


def _validate_relevance_weight(value: object, *, keyword: str) -> None:
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
        or not 0 <= value <= 1
    ):
        raise CatalogValidationError(
            f"invalid {RELEVANCE_WEIGHT_COLUMN} for {keyword}: {value}"
        )


def _optional_int(value: object, *, column: str, keyword: str) -> int | None:
    numeric_value = _optional_float(value, column=column, keyword=keyword)
    if numeric_value is None:
        return None
    if not numeric_value.is_integer():
        raise CatalogValidationError(f"invalid {column} for {keyword}: {value}")
    return int(numeric_value)


def load_keyword_catalog(path: Path) -> list[KeywordCatalogRow]:
    if not path.is_file():
        raise CatalogValidationError(
            "Keyword source workbook is missing: "
            f"{path}. Provision the approved June 2026 workbook at this path "
            "or pass it explicitly with market_monitor_import_keywords.py --workbook PATH."
        )
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        try:
            sheet = workbook[SHEET_NAME]
        except KeyError as exc:
            raise CatalogValidationError(f"missing sheet: {SHEET_NAME}") from exc

        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        index = {name: position for position, name in enumerate(headers)}
        missing = sorted(set(REQUIRED_COLUMNS) - set(index))
        if missing:
            raise CatalogValidationError(f"missing columns: {', '.join(missing)}")

        rows: list[KeywordCatalogRow] = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            keyword = normalize_keyword(values[index[KEYWORD_COLUMN]])
            if not keyword:
                continue

            source_level = _optional_text(values[index[RELEVANCE_COLUMN]])
            if source_level not in SOURCE_TO_LEVEL:
                raise CatalogValidationError(
                    f"invalid relevance for {keyword}: {source_level or ''}"
                )

            relevance_weight = _optional_float(
                values[index[RELEVANCE_WEIGHT_COLUMN]],
                column=RELEVANCE_WEIGHT_COLUMN,
                keyword=keyword,
            )
            if relevance_weight is None:
                raise CatalogValidationError(
                    f"missing {RELEVANCE_WEIGHT_COLUMN} for {keyword}"
                )

            rows.append(
                KeywordCatalogRow(
                    keyword=keyword,
                    translation=_optional_text(values[index[TRANSLATION_COLUMN]]),
                    relevance_level=SOURCE_TO_LEVEL[source_level],
                    relevance_weight=relevance_weight,
                    keyword_rank=_optional_int(
                        values[index["关键词排名"]], column="关键词排名", keyword=keyword
                    ),
                    search_volume=_optional_float(
                        values[index["月搜索量"]], column="月搜索量", keyword=keyword
                    ),
                    category_search_volume=_optional_float(
                        values[index["类目相关搜索量"]],
                        column="类目相关搜索量",
                        keyword=keyword,
                    ),
                    cpc=_optional_float(
                        values[index["CPC建议竞价($)"]],
                        column="CPC建议竞价($)",
                        keyword=keyword,
                    ),
                    cpc_range=_optional_text(values[index["建议竞价范围($)"]]),
                    click_conversion_rate=_optional_float(
                        values[index["点击转化率(均值)"]],
                        column="点击转化率(均值)",
                        keyword=keyword,
                    ),
                    competitive_difficulty=_optional_float(
                        values[index["竞争难度"]], column="竞争难度", keyword=keyword
                    ),
                    organic_scroll_rate=_optional_float(
                        values[index["自然位滚动率"]],
                        column="自然位滚动率",
                        keyword=keyword,
                    ),
                )
            )
    finally:
        workbook.close()

    validate_keyword_catalog(rows)
    return rows


def validate_keyword_catalog(
    rows: Iterable[KeywordCatalogRow],
    expected_counts: dict[str, int] | None = EXPECTED_KEYWORD_COUNTS,
) -> dict[str, int]:
    counts = {level: 0 for level in RELEVANCE_LEVELS}
    seen: set[str] = set()

    for row in rows:
        keyword = normalize_keyword(row.keyword)
        if not keyword:
            raise CatalogValidationError("empty keyword")
        if keyword in seen:
            raise CatalogValidationError(f"duplicate keyword: {keyword}")
        seen.add(keyword)

        if row.relevance_level not in counts:
            raise CatalogValidationError(f"invalid relevance: {row.relevance_level}")
        _validate_relevance_weight(row.relevance_weight, keyword=keyword)
        counts[row.relevance_level] += 1

    if expected_counts is not None and counts != expected_counts:
        raise CatalogValidationError(f"unexpected relevance counts: {counts}")
    return counts
